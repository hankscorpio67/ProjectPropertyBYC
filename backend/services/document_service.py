"""
Document ingestion: parse PDF/DOCX/XLSX/TXT files, chunk text, index into ChromaDB.
"""
import os
import json
import asyncio
from typing import List, Dict
import tiktoken

from ..config import config
from .vector_store import VectorStore

_tokenizer = tiktoken.get_encoding("cl100k_base")
CHUNK_SIZE = 800   # tokens per chunk
CHUNK_OVERLAP = 100  # token overlap between chunks


def _count_tokens(text: str) -> int:
    return len(_tokenizer.encode(text))


def _split_into_chunks(text: str, source: str = "") -> List[Dict]:
    """Split text into overlapping token-bounded chunks."""
    tokens = _tokenizer.encode(text)
    chunks = []
    start = 0
    while start < len(tokens):
        end = min(start + CHUNK_SIZE, len(tokens))
        chunk_tokens = tokens[start:end]
        chunk_text = _tokenizer.decode(chunk_tokens)
        chunks.append({"text": chunk_text, "source": source, "token_count": len(chunk_tokens)})
        if end == len(tokens):
            break
        start = end - CHUNK_OVERLAP
    return chunks


def _parse_pdf(file_path: str) -> List[Dict]:
    """Parse PDF, preserving page numbers. Uses pdfplumber for table support."""
    import pdfplumber
    chunks = []
    with pdfplumber.open(file_path) as pdf:
        for page_num, page in enumerate(pdf.pages, 1):
            text = page.extract_text() or ""
            # Also extract tables as structured text
            tables = page.extract_tables()
            for table in tables:
                rows = []
                for row in table:
                    row_text = " | ".join(str(cell or "") for cell in row)
                    rows.append(row_text)
                if rows:
                    text += "\n\nTable:\n" + "\n".join(rows)
            if text.strip():
                page_chunks = _split_into_chunks(text.strip(), source=f"page {page_num}")
                chunks.extend(page_chunks)
    return chunks


def _parse_docx(file_path: str) -> List[Dict]:
    """Parse DOCX file."""
    from docx import Document
    doc = Document(file_path)
    parts = []
    for para in doc.paragraphs:
        if para.text.strip():
            parts.append(para.text)
    # Extract tables
    for table in doc.tables:
        rows = []
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells)
            if row_text.strip():
                rows.append(row_text)
        if rows:
            parts.append("Table:\n" + "\n".join(rows))
    full_text = "\n\n".join(parts)
    return _split_into_chunks(full_text, source="document")


def _parse_xlsx(file_path: str) -> List[Dict]:
    """Parse Excel spreadsheet, converting each sheet to structured text."""
    import openpyxl
    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    all_chunks = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = []
        for row in ws.iter_rows(values_only=True):
            row_vals = [str(v) if v is not None else "" for v in row]
            if any(v.strip() for v in row_vals):
                rows.append(" | ".join(row_vals))
        if rows:
            sheet_text = f"Sheet: {sheet_name}\n" + "\n".join(rows)
            sheet_chunks = _split_into_chunks(sheet_text, source=f"sheet:{sheet_name}")
            all_chunks.extend(sheet_chunks)
    wb.close()
    return all_chunks


def _parse_txt(file_path: str) -> List[Dict]:
    """Parse plain text file."""
    with open(file_path, "r", encoding="utf-8", errors="replace") as f:
        text = f.read()
    return _split_into_chunks(text, source="text file")


class DocumentService:
    def __init__(self):
        self._vector_store = VectorStore()

    async def ingest(self, doc_id: str, project_id: str, file_path: str, file_type: str) -> int:
        """Parse file, chunk text, store in vector DB. Returns chunk count."""
        loop = asyncio.get_event_loop()
        chunks = await loop.run_in_executor(
            None, self._parse_file, file_path, file_type
        )
        if not chunks:
            return 0

        # Prepare for vector store
        texts = [c["text"] for c in chunks]
        ids = [f"{doc_id}_chunk_{i}" for i in range(len(chunks))]
        metadatas = [
            {
                "doc_id": doc_id,
                "project_id": project_id,
                "source": c.get("source", ""),
                "chunk_index": i,
            }
            for i, c in enumerate(chunks)
        ]

        await loop.run_in_executor(
            None, self._vector_store.add, project_id, texts, ids, metadatas
        )
        return len(chunks)

    def _parse_file(self, file_path: str, file_type: str) -> List[Dict]:
        parsers = {
            "pdf": _parse_pdf,
            "docx": _parse_docx,
            "doc": _parse_docx,
            "xlsx": _parse_xlsx,
            "xls": _parse_xlsx,
            "txt": _parse_txt,
        }
        parser = parsers.get(file_type)
        if not parser:
            return []
        return parser(file_path)

    def delete_document(self, doc_id: str, project_id: str):
        """Remove all chunks for a document from the vector store."""
        self._vector_store.delete_by_doc(project_id, doc_id)

    async def search(self, project_id: str, query: str, n_results: int = None) -> List[Dict]:
        """Semantic search for relevant chunks."""
        k = n_results or config.RETRIEVAL_CHUNKS
        loop = asyncio.get_event_loop()
        results = await loop.run_in_executor(
            None, self._vector_store.query, project_id, query, k
        )
        return results
